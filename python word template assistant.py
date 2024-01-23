from docx import Document
import os

def create_invoice(template_path, output_path, data):
    # Load the template
    doc = Document(template_path)

    # Replace placeholders with actual data
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)

    # Save the new document
    doc.save(output_path)

# Define the data to replace placeholders
invoice_data = {
    '{CUSTOMER_NAME}': 'John Doe',
    '{INVOICE_NUMBER}': '12345',
    '{INVOICE_DATE}': '2024-01-23',
    # Add more placeholders and their replacements here
}

# Specify the paths
template_path = 'path/to/your/InvoiceTemplate.docx'
output_path = 'path/to/output/Invoice_12345.docx'

# Create the invoice
create_invoice(template_path, output_path, invoice_data)

print("Invoice created successfully.")
